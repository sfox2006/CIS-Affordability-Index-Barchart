"""Check release discovery and safe refresh behavior independently of live ABS."""
import importlib.util
import unittest
from datetime import datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import openpyxl

spec = importlib.util.spec_from_file_location("update", Path(__file__).resolve().parents[1] / "scripts/update_abs_data.py")
update = importlib.util.module_from_spec(spec)
spec.loader.exec_module(update)


class RefreshTests(unittest.TestCase):
    def test_release_discovery(self):
        html = b'<a href="/release/6401018.xlsx">Table 18</a><a href="/release/6401017.xlsx">Other</a>'
        with patch.object(update, "download", return_value=html):
            self.assertEqual(update.latest_file("https://www.abs.gov.au/latest", "6401018.xlsx"),
                             "https://www.abs.gov.au/release/6401018.xlsx")

    def test_missing_release_fails(self):
        with patch.object(update, "download", return_value=b"Unavailable"):
            with self.assertRaises(ValueError):
                update.latest_file("https://www.abs.gov.au/latest", "634501.xlsx")

    def test_workbook_validation(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Data1"
        sheet.cell(2, 2, "Index Numbers")
        sheet.cell(4, 2, "INDEX")
        sheet.cell(10, 2, "A1")
        sheet.cell(11, 1, datetime(2026, 6, 1))
        sheet.cell(11, 2, 100)
        buffer = BytesIO()
        book.save(buffer)
        self.assertEqual(update.workbook_series(buffer.getvalue(), ["A1"])["A1"],
                         [{"date": "2026-06-01", "value": 100}])
        with self.assertRaises(ValueError):
            update.workbook_series(buffer.getvalue(), ["MISSING"])
        sheet.cell(11, 2, 0)
        buffer = BytesIO()
        book.save(buffer)
        with self.assertRaises(ValueError):
            update.workbook_series(buffer.getvalue(), ["A1"])

    def test_new_quarter_revision_and_unchanged(self):
        current = update.read_bundle(update.ROOT / "data.js")
        cpi = {s["seriesId"]: s["observations"] for s in current["CPI_DATA"]["series"]}
        sid = current["WPI_METADATA"]["seriesId"]
        wpi = {sid: current["WPI_DATA"]}
        self.assertEqual(update.rebuild(current, cpi, wpi), current)
        wpi[sid] = current["WPI_DATA"] + [{"date": "2099-09-01", "value": 200}]
        result = update.rebuild(current, cpi, wpi)
        self.assertEqual(update.shared_end(result), update.shared_end(current))
        wpi[sid] = [dict(p) for p in current["WPI_DATA"]]
        wpi[sid][-1]["value"] += 1
        self.assertNotEqual(update.rebuild(current, cpi, wpi), current)

    def test_missing_history_fails(self):
        with self.assertRaises(ValueError):
            update.preserve_history([{"date": "2026-06-01"}], [{"date": "2026-09-01"}])


if __name__ == "__main__":
    unittest.main()
